"""
steps_summary.py — ESC Pipeline Conversion Rate Analysis

Loads method summary JSONs (or raw results JSONs as fallback) and prints
a formatted table showing the step-by-step conversion funnel for each
model x benchmark combination. Also saves results to Excel.

MMSafetyBench: loads summary1/2/3 files and SUMS counts (not averages)
               so that N reflects the real total samples across all subsets.

Usage:
    python steps_summary.py
"""

import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SUMMARY_DIR  = "/workspace/results/method_summary"
OUTPUT_EXCEL = "/workspace/results/method_summary/esc_funnel_results.xlsx"

MODELS = ["llava", "qwen"]

BENCHMARKS_DISPLAY_ORDER = [
    "vlsafe", "mmsafety", "blink", "pope", "rwqa",
    "hallusion", "mmvp", "mathvista", "mmstar", "mmvet",
]

MODEL_DISPLAY = {
    "llava": "LLaVA-1.5-7B",
    "qwen":  "Qwen2-VL-7B",
}

BENCH_DISPLAY = {
    "vlsafe":    "VLSafe",
    "mmsafety":  "MMSafetyBench (sum)",
    "blink":     "BLINK",
    "pope":      "POPE",
    "rwqa":      "RWQA",
    "hallusion": "HallusionBench",
    "mmvp":      "MMVP",
    "mathvista": "MathVista",
    "mmstar":    "MMStar",
    "mmvet":     "MM-Vet",
}

# ==============================================================================
# LOADERS
# ==============================================================================

def load_from_summary(path):
    with open(path) as f:
        d = json.load(f)
    return {
        "total":               d["total_samples"],
        "step2_safe":          d["step2_judged_safe"],
        "step2_unsafe":        d["step2_judged_unsafe"],
        "step5_regenerated":   d["step5_chose_regenerated"],
        "step5_kept_original": d["step5_kept_original"],
    }


def load_from_results(path):
    with open(path) as f:
        data = json.load(f)
    total        = len(data)
    step2_safe   = sum(1 for s in data if s.get("judge_is_safe", False))
    step2_unsafe = total - step2_safe
    step5_regen  = sum(1 for s in data if s.get("was_regenerated", False))
    return {
        "total":               total,
        "step2_safe":          step2_safe,
        "step2_unsafe":        step2_unsafe,
        "step5_regenerated":   step5_regen,
        "step5_kept_original": step2_unsafe - step5_regen,
    }


def load_single(model, benchmark):
    spath = os.path.join(SUMMARY_DIR, f"method_{model}_{benchmark}_summary.json")
    rpath = os.path.join(SUMMARY_DIR, f"method_{model}_{benchmark}_results.json")
    if os.path.exists(spath):
        print(f"  [summary] {os.path.basename(spath)}")
        return load_from_summary(spath)
    elif os.path.exists(rpath):
        print(f"  [results] {os.path.basename(rpath)}")
        return load_from_results(rpath)
    else:
        print(f"  [MISSING] {model} / {benchmark}")
        return None


def load_mmsafety_summed(model):
    """Sum counts across 3 subsets — N is real total, not a decimal average."""
    stats_list = []
    for i in range(1, 4):
        spath = os.path.join(SUMMARY_DIR, f"method_{model}_mmsafety_summary{i}.json")
        rpath = os.path.join(SUMMARY_DIR, f"method_{model}_mmsafety_results{i}.json")
        if os.path.exists(spath):
            print(f"  [summary] {os.path.basename(spath)}")
            stats_list.append(load_from_summary(spath))
        elif os.path.exists(rpath):
            print(f"  [results] {os.path.basename(rpath)}")
            stats_list.append(load_from_results(rpath))
        else:
            print(f"  [MISSING] method_{model}_mmsafety_summary{i}.json")

    if not stats_list:
        return None

    summed = {
        key: sum(s[key] for s in stats_list)
        for key in ["total", "step2_safe", "step2_unsafe",
                    "step5_regenerated", "step5_kept_original"]
    }
    print(f"  -> Summed {len(stats_list)} subsets  (N={summed['total']})")
    return summed


def compute_rates(stats):
    total  = stats["total"]
    unsafe = stats["step2_unsafe"]
    regen  = stats["step5_regenerated"]
    return {
        **stats,
        "flag_rate":            unsafe / total  if total  > 0 else 0,
        "regen_accept":         regen  / unsafe if unsafe > 0 else 0,
        "overall_intervention": regen  / total  if total  > 0 else 0,
    }

# ==============================================================================
# CONSOLE TABLE
# ==============================================================================

def fmt_int(v): return str(int(round(v)))
def fmt_pct(v): return f"{v * 100:.1f}%"


def print_table(available):
    W = dict(model=15, bench=22, n=7, safe=7, flag=16, regen=18, kept=10, inter=14)
    hdr = (
        "{:<{model}} {:<{bench}} {:>{n}} {:>{safe}} {:>{flag}} {:>{regen}} {:>{kept}} {:>{inter}}"
    ).format(
        "Model", "Benchmark", "N", "Safe",
        "Flagged (S2)", "Regenerated (S5)", "Kept Orig", "Intervention",
        **W
    )
    sep  = "=" * len(hdr)
    thin = "-" * len(hdr)
    print(f"\n{sep}\nESC PIPELINE CONVERSION FUNNEL\n{sep}\n{hdr}\n{sep}")

    models_found = [m for m in MODELS if any(m == k[0] for k in available)]
    for mi, model in enumerate(models_found):
        rows = sorted(
            [(m, b) for (m, b) in available if m == model],
            key=lambda x: (BENCHMARKS_DISPLAY_ORDER.index(x[1])
                           if x[1] in BENCHMARKS_DISPLAY_ORDER else 999)
        )
        for ji, (m, b) in enumerate(rows):
            r          = available[(m, b)]
            model_lbl  = MODEL_DISPLAY.get(m, m) if ji == 0 else ""
            flag_str   = "{} ({})".format(fmt_int(r["step2_unsafe"]),  fmt_pct(r["flag_rate"]))
            regen_str  = "{} ({})".format(fmt_int(r["step5_regenerated"]), fmt_pct(r["regen_accept"]))
            print(
                "{:<{model}} {:<{bench}} {:>{n}} {:>{safe}} {:>{flag}} {:>{regen}} {:>{kept}} {:>{inter}}".format(
                    model_lbl, BENCH_DISPLAY.get(b, b),
                    fmt_int(r["total"]), fmt_int(r["step2_safe"]),
                    flag_str, regen_str,
                    fmt_int(r["step5_kept_original"]), fmt_pct(r["overall_intervention"]),
                    **W
                )
            )
        if mi < len(models_found) - 1:
            print(thin)
    print(sep)

# ==============================================================================
# EXCEL EXPORT
# ==============================================================================

def _border():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(hex_):
    return PatternFill("solid", start_color=hex_, fgColor=hex_)


def save_excel(available, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "ESC Funnel"

    border = _border()
    center = Alignment(horizontal="center", vertical="center")
    left   = Alignment(horizontal="left",   vertical="center")

    # ── Title ──
    ws.merge_cells("A1:I1")
    ws["A1"] = "ESC Pipeline — Conversion Funnel"
    ws["A1"].font      = Font(name="Arial", bold=True, size=13, color="FFFFFF")
    ws["A1"].fill      = _fill("1F3864")
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 26

    # ── Column headers ──
    COL_HEADERS = [
        "Model", "Benchmark", "N (Total)",
        "Safe\n(Step 2)", "Flagged\n(Step 2)", "Flag Rate",
        "Regenerated\n(Step 5)", "Regen Accept\nRate", "Overall\nIntervention",
    ]
    for ci, h in enumerate(COL_HEADERS, 1):
        c = ws.cell(row=2, column=ci, value=h)
        c.font      = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill      = _fill("2F5496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = border
    ws.row_dimensions[2].height = 34

    # ── Column widths ──
    for i, w in enumerate([17, 23, 10, 12, 12, 11, 16, 15, 17], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Data ──
    row = 3
    models_found = [m for m in MODELS if any(m == k[0] for k in available)]

    for mi, model in enumerate(models_found):
        bench_rows = sorted(
            [(m, b) for (m, b) in available if m == model],
            key=lambda x: (BENCHMARKS_DISPLAY_ORDER.index(x[1])
                           if x[1] in BENCHMARKS_DISPLAY_ORDER else 999)
        )
        model_start = row

        for ji, (m, b) in enumerate(bench_rows):
            r     = available[(m, b)]
            bg    = "FFFFFF" if ji % 2 == 0 else "EBF3FB"
            rfill = _fill(bg)
            bfont = Font(name="Arial", size=10)

            values = [
                "",                                     # col 1 — set via merge below
                BENCH_DISPLAY.get(b, b),                # col 2
                int(round(r["total"])),                 # col 3
                int(round(r["step2_safe"])),            # col 4
                int(round(r["step2_unsafe"])),          # col 5
                r["flag_rate"],                         # col 6
                int(round(r["step5_regenerated"])),     # col 7
                r["regen_accept"],                      # col 8
                r["overall_intervention"],              # col 9
            ]

            for ci, val in enumerate(values, 1):
                c = ws.cell(row=row, column=ci, value=val)
                c.font   = bfont
                c.fill   = rfill
                c.border = border
                if ci in (3, 4, 5, 7):
                    c.alignment     = center
                    c.number_format = "#,##0"
                elif ci in (6, 8, 9):
                    c.alignment     = center
                    c.number_format = "0.0%"
                else:
                    c.alignment = left
            row += 1

        # Merge model column and style it
        if len(bench_rows) > 1:
            ws.merge_cells(start_row=model_start, start_column=1,
                           end_row=row - 1,       end_column=1)
        mc           = ws.cell(row=model_start, column=1)
        mc.value     = MODEL_DISPLAY.get(model, model)
        mc.font      = Font(name="Arial", bold=True, size=10)
        mc.fill      = _fill("D6E4F0")
        mc.alignment = Alignment(horizontal="center", vertical="center")
        mc.border    = border

        # Separator row between model blocks
        if mi < len(models_found) - 1:
            for ci in range(1, len(COL_HEADERS) + 1):
                c      = ws.cell(row=row, column=ci, value="")
                c.fill = _fill("C9D9EE")
                c.border = border
            ws.row_dimensions[row].height = 5
            row += 1

    ws.freeze_panes = "A3"
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print(f"\n💾 Excel saved -> {path}")

# ==============================================================================
# MAIN
# ==============================================================================

def main():
    print("Loading ESC pipeline funnel data...\n")
    available = {}

    for model in MODELS:
        print(f"{MODEL_DISPLAY.get(model, model)}:")
        for bench in BENCHMARKS_DISPLAY_ORDER:
            if bench == "mmsafety":
                stats = load_mmsafety_summed(model)
            else:
                stats = load_single(model, bench)
            if stats is not None:
                available[(model, bench)] = compute_rates(stats)
        print()

    if not available:
        print("No data found! Check SUMMARY_DIR path.")
        return

    print_table(available)
    save_excel(available, OUTPUT_EXCEL)


if __name__ == "__main__":
    main()